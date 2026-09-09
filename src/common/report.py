"""수집 결과(data/*.jsonl)를 Excel 워크북으로 묶어 이메일 발송.

- 데이터셋별로 시트 1개 (trends, serp, llm_mentions, news_volume, ...)
- 'summary' 시트에 최신 핵심 지표 요약
- SMTP(기본 Gmail)로 첨부 발송
"""
from __future__ import annotations
import json, io, smtplib, datetime as dt
from email.message import EmailMessage
from pathlib import Path
from .config import ROOT, SETTINGS, env, today

DATA_DIR = ROOT / SETTINGS["storage"]["local_dir"]


def _read_jsonl(path: Path) -> list[dict]:
    rows = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def build_workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    datasets = sorted(DATA_DIR.glob("*.jsonl"))

    summary = wb.create_sheet("summary")
    summary.append(["dataset", "총 행수", "마지막 수집(UTC)"])

    for path in datasets:
        rows = _read_jsonl(path)
        if not rows:
            continue
        name = path.stem[:31]
        ws = wb.create_sheet(name)
        header = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(header)
        for r in rows:
            ws.append([_cell(r.get(h, "")) for h in header])
        last = max((r.get("_ingested_at", "") for r in rows), default="")
        summary.append([name, len(rows), last])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return v


def email_report(subject: str | None = None, body: str = "") -> None:
    to = env("REPORT_EMAIL_TO")
    user = env("SMTP_USER")
    pw = env("SMTP_PASS")
    if not (to and user and pw):
        print("[report] SMTP_USER/SMTP_PASS/REPORT_EMAIL_TO 미설정 → 발송 생략")
        return

    xlsx = build_workbook()
    msg = EmailMessage()
    msg["Subject"] = subject or f"[애플 폴더블 모니터링] {today()} 리포트"
    msg["From"] = user
    msg["To"] = to
    msg.set_content(body or "첨부된 Excel 워크북에 최신 수집 데이터가 담겨 있습니다.")
    msg.add_attachment(
        xlsx, maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"apple_foldable_monitor_{today()}.xlsx",
    )

    host = env("SMTP_HOST", "smtp.gmail.com")
    port = int(env("SMTP_PORT", "465"))
    with smtplib.SMTP_SSL(host, port) as s:
        s.login(user, pw)
        s.send_message(msg)
    print(f"[report] 발송 완료 → {to}")


if __name__ == "__main__":
    email_report()
